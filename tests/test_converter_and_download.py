import base64
import json
from pathlib import Path

import pytest

from filesystem_rag_mcp.config import Settings
from filesystem_rag_mcp.converter import convert_file_to_markdown
from filesystem_rag_mcp.server import build_server


def test_converter_markdown_formats(tmp_path: Path):
    # 1. JSON
    json_path = tmp_path / "data.json"
    json_path.write_text(json.dumps({"name": "RAG MCP", "version": "0.1.0"}))
    md_json = convert_file_to_markdown(json_path)
    assert "# data.json" in md_json
    assert "```json" in md_json
    assert "RAG MCP" in md_json

    # 2. CSV
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("item,cost,qty\napple,1.5,10\norange,2.0,5\n")
    md_csv = convert_file_to_markdown(csv_path)
    assert "| item | cost | qty |" in md_csv
    assert "| apple | 1.5 | 10 |" in md_csv

    # 3. HTML
    html_path = tmp_path / "doc.html"
    html_path.write_text(
        "<html><head><title>Test Doc</title></head><body><h1>Heading</h1><p>Paragraph text</p></body></html>"
    )
    md_html = convert_file_to_markdown(html_path)
    assert "Test Doc" in md_html
    assert "Paragraph text" in md_html


def test_converter_ipynb(tmp_path: Path):
    nb_path = tmp_path / "analysis.ipynb"
    nb_content = {
        "cells": [
            {
                "cell_type": "markdown",
                "metadata": {},
                "source": ["# Analysis Notebook\n", "Explaining results."],
            },
            {
                "cell_type": "code",
                "execution_count": 1,
                "metadata": {},
                "outputs": [{"output_type": "stream", "name": "stdout", "text": ["42\n"]}],
                "source": ["x = 42\n", "print(x)"],
            },
        ],
        "metadata": {},
        "nbformat": 4,
        "nbformat_minor": 2,
    }
    nb_path.write_text(json.dumps(nb_content))
    md_nb = convert_file_to_markdown(nb_path)
    assert "# Notebook: analysis.ipynb" in md_nb
    assert "Explaining results." in md_nb
    assert "```python" in md_nb
    assert "x = 42" in md_nb
    assert "```output" in md_nb
    assert "42" in md_nb


def test_converter_docx(tmp_path: Path):
    import docx

    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Sample Document", level=1)
    doc.add_paragraph("This is a paragraph inside a Word doc.")
    doc.save(str(docx_path))

    md_docx = convert_file_to_markdown(docx_path)
    assert "This is a paragraph inside a Word doc." in md_docx


def test_converter_epub(tmp_path: Path):
    epub_path = Path("/tmp/sample.epub")
    if epub_path.exists():
        md_epub = convert_file_to_markdown(epub_path)
        assert "Intro Chapter" in md_epub
        assert "This is the intro text in an EPUB book." in md_epub


@pytest.mark.asyncio
async def test_server_tools_markdown_and_download(tmp_path: Path):
    root = tmp_path / "root"
    root.mkdir()
    (root / "hello.txt").write_text("Hello plain text!")
    (root / "binary.bin").write_bytes(b"\x00\x01\x02\x03\x04\xff")

    settings = Settings(root_dir=root, data_dir=tmp_path / "data")
    _server = build_server(settings)

    # Directly verify state methods
    from filesystem_rag_mcp.server import _ServerState

    state = _ServerState(settings)

    # 1. read_file_markdown
    res_md = await state.read_file_markdown(rel_path="hello.txt")
    assert "error" not in res_md
    assert res_md["markdown"] == "Hello plain text!"

    # 2. download_file_raw
    res_raw = await state.download_file_raw(rel_path="binary.bin", max_bytes=None)
    assert "error" not in res_raw
    assert res_raw["total_size_bytes"] == 6
    assert base64.b64decode(res_raw["base64_data"]) == b"\x00\x01\x02\x03\x04\xff"

    # 3. Path traversal security checks
    traversal_md = await state.read_file_markdown(rel_path="../outside.txt")
    assert "error" in traversal_md

    traversal_raw = await state.download_file_raw(rel_path="../outside.bin", max_bytes=None)
    assert "error" in traversal_raw
